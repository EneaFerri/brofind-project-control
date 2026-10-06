"""Fail-closed importer for the frozen B1 economic-control workbook."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from typing import Any
import zipfile

import openpyxl
from sqlalchemy import delete
from sqlalchemy.orm import Session

from .models import Invoice, Project, ProjectControlLine, ProjectControlSnapshot


class B1ImportError(ValueError):
    pass


SHEETS = ("CONTROL_SNAPSHOT", "CONTROL_LINES", "INVOICES")
SIGNATURE = frozenset(("CONTROL_SNAPSHOT", "CONTROL_LINES"))
HEADERS = {
    "CONTROL_SNAPSHOT": ("Project_ID", "Snapshot_Date", "Source_Label", "Baseline_Budget", "Operational_Budget", "Orders_JMRP", "Orders_Extract", "Received", "Open_Order_Commitment", "Labor", "Overhead", "Note"),
    "CONTROL_LINES": ("Project_ID", "Snapshot_Date", "VDA_Code", "VDA_Name", "Analysis_Center", "Baseline_Budget", "Operational_Budget", "Ordered", "Received", "Open_Order_Commitment", "Labor", "Overhead", "Purchases_Closed", "Source_Ref"),
    "INVOICES": ("Project_ID", "Invoice_Number", "Milestone", "Issue_Date", "Due_Date", "Amount", "Paid_Amount"),
}
VDA_CODE = re.compile(r"C\d{6}\Z")
EPSILON = Decimal("0.01")


def _error(message: str) -> B1ImportError:
    return B1ImportError(f"B1: {message}")


def is_b1_workbook_candidate(content: bytes) -> bool:
    """Uses sheet signature only: invalid B1 must not fall back to generic."""
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception:
        return False
    return bool(SIGNATURE.intersection(workbook.sheetnames))


def _workbook(content: bytes):
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if any(name.lower().endswith("vbaproject.bin") for name in archive.namelist()):
                raise _error("macro non ammessa")
    except B1ImportError:
        raise
    except (OSError, zipfile.BadZipFile) as exc:
        raise _error("file XLSX non leggibile") from exc
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=False, keep_vba=False)
    except Exception as exc:
        raise _error("workbook XLSX non leggibile") from exc
    if tuple(workbook.sheetnames) != SHEETS:
        raise _error(f"fogli non conformi: attesi esattamente {', '.join(SHEETS)}")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    raise _error(f"formula non ammessa ({sheet.title}!{cell.coordinate})")
    return workbook


def _rows(sheet, expected: tuple[str, ...]) -> list[dict[str, Any]]:
    header = tuple(sheet.cell(1, col).value for col in range(1, sheet.max_column + 1))
    if header != expected:
        raise _error(f"header non conforme nel foglio {sheet.title}")
    output = []
    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row, col).value for col in range(1, len(expected) + 1)]
        if any(value is not None for value in values):
            output.append(dict(zip(expected, values)))
    return output


def _text(value: Any, field: str, optional: bool = False) -> str | None:
    if value is None or value == "":
        if optional:
            return None
        raise _error(f"{field} deve essere testo non vuoto")
    if not isinstance(value, str) or not value.strip():
        raise _error(f"{field} deve essere testo non vuoto")
    return value.strip()


def _date(value: Any, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            parsed = date.fromisoformat(value)
        except ValueError as exc:
            raise _error(f"{field} deve essere data ISO YYYY-MM-DD") from exc
        if parsed.isoformat() == value:
            return parsed
    raise _error(f"{field} deve essere data ISO YYYY-MM-DD")


def _amount(value: Any, field: str) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise _error(f"{field} deve essere un numero")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise _error(f"{field} deve essere un numero valido") from exc
    if not result.is_finite() or result < 0:
        raise _error(f"{field} deve essere finito e non negativo")
    return result


def _boolean(value: Any, field: str) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip().upper() in ("TRUE", "FALSE"):
        return value.strip().upper() == "TRUE"
    raise _error(f"{field} deve essere TRUE, FALSE o vuoto")


def _equal(left: Decimal, right: Decimal, label: str) -> None:
    if abs(left - right) > EPSILON:
        raise _error(f"quadratura non valida per {label}: {left} != {right}")


@dataclass(frozen=True)
class ParsedB1:
    snapshot: dict[str, Any]
    lines: tuple[dict[str, Any], ...]
    invoices: tuple[dict[str, Any], ...]


def _parse(content: bytes) -> ParsedB1:
    workbook = _workbook(content)
    snapshot_rows = _rows(workbook["CONTROL_SNAPSHOT"], HEADERS["CONTROL_SNAPSHOT"])
    line_rows = _rows(workbook["CONTROL_LINES"], HEADERS["CONTROL_LINES"])
    invoice_rows = _rows(workbook["INVOICES"], HEADERS["INVOICES"])
    if len(snapshot_rows) != 1 or not line_rows or not invoice_rows:
        raise _error("CONTROL_SNAPSHOT richiede una riga; CONTROL_LINES e INVOICES almeno una")
    raw = snapshot_rows[0]
    snapshot = {
        "project_code": _text(raw["Project_ID"], "CONTROL_SNAPSHOT.Project_ID"),
        "snapshot_date": _date(raw["Snapshot_Date"], "CONTROL_SNAPSHOT.Snapshot_Date"),
        "source_label": _text(raw["Source_Label"], "CONTROL_SNAPSHOT.Source_Label"),
        "note": _text(raw["Note"], "CONTROL_SNAPSHOT.Note", optional=True),
    }
    for key in ("Baseline_Budget", "Operational_Budget", "Orders_JMRP", "Orders_Extract", "Received", "Open_Order_Commitment", "Labor", "Overhead"):
        snapshot[key] = _amount(raw[key], f"CONTROL_SNAPSHOT.{key}")
    lines, keys, names = [], set(), {}
    for position, raw in enumerate(line_rows, 2):
        prefix = f"CONTROL_LINES riga {position}"
        line = {"project_code": _text(raw["Project_ID"], f"{prefix}.Project_ID"), "snapshot_date": _date(raw["Snapshot_Date"], f"{prefix}.Snapshot_Date"), "vda_code": _text(raw["VDA_Code"], f"{prefix}.VDA_Code"), "vda_name": _text(raw["VDA_Name"], f"{prefix}.VDA_Name"), "analysis_center": _text(raw["Analysis_Center"], f"{prefix}.Analysis_Center"), "source_ref": _text(raw["Source_Ref"], f"{prefix}.Source_Ref"), "purchases_closed": _boolean(raw["Purchases_Closed"], f"{prefix}.Purchases_Closed")}
        for key in ("Baseline_Budget", "Operational_Budget", "Ordered", "Received", "Open_Order_Commitment", "Labor", "Overhead"):
            line[key] = _amount(raw[key], f"{prefix}.{key}")
        if not VDA_CODE.fullmatch(line["vda_code"]):
            raise _error(f"{prefix}.VDA_Code non conforme")
        if line["project_code"] != snapshot["project_code"] or line["snapshot_date"] != snapshot["snapshot_date"]:
            raise _error(f"{prefix} non appartiene allo snapshot")
        key = (line["vda_code"], line["analysis_center"])
        if key in keys:
            raise _error(f"duplicata la coppia VDA/CDA {key[0]}/{key[1]}")
        keys.add(key)
        if line["vda_code"] in names and names[line["vda_code"]] != line["vda_name"]:
            raise _error(f"VDA_Name incoerente per {line['vda_code']}")
        names[line["vda_code"]] = line["vda_name"]
        lines.append(line)
    invoices, invoice_numbers = [], set()
    for position, raw in enumerate(invoice_rows, 2):
        prefix = f"INVOICES riga {position}"
        invoice = {"project_code": _text(raw["Project_ID"], f"{prefix}.Project_ID"), "number": _text(raw["Invoice_Number"], f"{prefix}.Invoice_Number"), "milestone": _text(raw["Milestone"], f"{prefix}.Milestone"), "issue_date": _date(raw["Issue_Date"], f"{prefix}.Issue_Date"), "due_date": _date(raw["Due_Date"], f"{prefix}.Due_Date"), "amount": _amount(raw["Amount"], f"{prefix}.Amount"), "paid_amount": _amount(raw["Paid_Amount"], f"{prefix}.Paid_Amount")}
        if invoice["project_code"] != snapshot["project_code"] or invoice["due_date"] < invoice["issue_date"] or invoice["paid_amount"] > invoice["amount"]:
            raise _error(f"{prefix} non coerente")
        if invoice["number"] in invoice_numbers:
            raise _error(f"Invoice_Number duplicato: {invoice['number']}")
        invoice_numbers.add(invoice["number"])
        invoices.append(invoice)
    _equal(snapshot["Orders_JMRP"], snapshot["Orders_Extract"], "Orders_JMRP / Orders_Extract")
    _equal(snapshot["Orders_Extract"], snapshot["Received"] + snapshot["Open_Order_Commitment"], "ordini / ricevuto + impegni")
    for snapshot_key, line_key in (("Baseline_Budget", "Baseline_Budget"), ("Operational_Budget", "Operational_Budget"), ("Orders_Extract", "Ordered"), ("Received", "Received"), ("Open_Order_Commitment", "Open_Order_Commitment"), ("Labor", "Labor"), ("Overhead", "Overhead")):
        _equal(snapshot[snapshot_key], sum(line[line_key] for line in lines), f"snapshot / CONTROL_LINES {snapshot_key}")
    return ParsedB1(snapshot, tuple(lines), tuple(invoices))


def _one(records: list[Any], key: str):
    if len(records) > 1:
        raise _error(f"duplicati gia presenti nel DB per {key}; import annullato")
    return records[0] if records else None


def _upsert_invoice(db: Session, project: Project, source: dict[str, Any]) -> bool:
    invoice = _one(db.query(Invoice).filter(Invoice.project_id == project.id, Invoice.number == source["number"]).all(), f"fattura {project.code}/{source['number']}")
    created = invoice is None
    if created:
        invoice = Invoice(project_id=project.id, number=source["number"])
        db.add(invoice)
    invoice.milestone_name, invoice.issue_date, invoice.due_date = source["milestone"], source["issue_date"], source["due_date"]
    invoice.amount, invoice.paid_amount = float(source["amount"]), float(source["paid_amount"])
    return created


def import_b1_control_workbook(content: bytes, db: Session, *, expected_project_code: str | None = None, commit: bool = True) -> dict[str, int]:
    """Validate first, then persist B1 in one rollback-safe transaction."""
    parsed = _parse(content)
    if expected_project_code is not None and parsed.snapshot["project_code"] != expected_project_code:
        raise _error("il Project_ID del workbook non corrisponde alla commessa selezionata")
    stats = {"control_snapshots": 0, "control_lines": 0, "invoices": 0}
    try:
        with db.begin_nested():
            source = parsed.snapshot
            project = db.query(Project).filter(Project.code == source["project_code"]).one_or_none()
            if project is None:
                raise _error(f"progetto inesistente: {source['project_code']}")
            captured = datetime.combine(source["snapshot_date"], datetime.min.time())
            snapshot = _one(db.query(ProjectControlSnapshot).filter(ProjectControlSnapshot.project_id == project.id, ProjectControlSnapshot.captured_at == captured, ProjectControlSnapshot.source_label == source["source_label"]).all(), "snapshot")
            for invoice in parsed.invoices:
                _one(db.query(Invoice).filter(Invoice.project_id == project.id, Invoice.number == invoice["number"]).all(), f"fattura {project.code}/{invoice['number']}")
            if snapshot is None:
                snapshot = ProjectControlSnapshot(project_id=project.id, captured_at=captured, source_label=source["source_label"])
                db.add(snapshot)
                stats["control_snapshots"] = 1
            for field, value in (("baseline_budget", source["Baseline_Budget"]), ("operational_budget", source["Operational_Budget"]), ("orders_jmrp", source["Orders_JMRP"]), ("orders_extract", source["Orders_Extract"]), ("received", source["Received"]), ("open_order_commitment", source["Open_Order_Commitment"]), ("labor", source["Labor"]), ("overhead", source["Overhead"])):
                setattr(snapshot, field, float(value))
            snapshot.note = source["note"]
            db.flush()
            db.execute(delete(ProjectControlLine).where(ProjectControlLine.snapshot_id == snapshot.id))
            for line in parsed.lines:
                db.add(ProjectControlLine(snapshot_id=snapshot.id, vda_code=line["vda_code"], vda_name=line["vda_name"], analysis_center=line["analysis_center"], baseline_budget=float(line["Baseline_Budget"]), operational_budget=float(line["Operational_Budget"]), ordered=float(line["Ordered"]), received=float(line["Received"]), open_order_commitment=float(line["Open_Order_Commitment"]), labor=float(line["Labor"]), overhead=float(line["Overhead"]), purchases_closed=line["purchases_closed"], source_ref=line["source_ref"]))
            stats["control_lines"] = len(parsed.lines)
            db.flush()
            for invoice in parsed.invoices:
                stats["invoices"] += int(_upsert_invoice(db, project, invoice))
            db.flush()
        if commit:
            db.commit()
    except Exception:
        db.rollback()
        raise
    return stats