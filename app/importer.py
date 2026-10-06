from datetime import datetime
from io import BytesIO

import openpyxl
from sqlalchemy.orm import Session

from .models import (
    Action,
    BudgetLine,
    ChangeRequest,
    Communication,
    CustomerOrder,
    Invoice,
    Milestone,
    Payment,
    Project,
    ProjectControlLine,
    ProjectControlSnapshot,
)

SHEETS = [
    "PROJECTS",
    "MILESTONES",
    "BUDGET",
    "COSTS",
    "INVOICES",
    "PAYMENTS",
    "ACTIONS",
    "COMMUNICATIONS",
    "CUSTOMER_ORDERS",
    "CHANGES",
    "CONTROL_SNAPSHOT",
    "CONTROL_LINES",
]


def _d(v):
    if v is None or v == "":
        return None
    if hasattr(v, "date") and not isinstance(v, str):
        return v.date() if isinstance(v, datetime) else v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(v, fmt).date()
            except ValueError:
                pass
    return v


def _rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(x).strip() if x is not None else "" for x in rows[0]]
    return [dict(zip(headers, row, strict=False)) for row in rows[1:] if any(v is not None for v in row)]


def import_workbook(content: bytes, db: Session):
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    if "PROJECTS" not in wb.sheetnames:
        raise ValueError("Foglio PROJECTS mancante")
    stats = {
        "projects": 0,
        "milestones": 0,
        "budget": 0,
        "invoices": 0,
        "payments": 0,
        "actions": 0,
        "communications": 0,
        "customer_orders": 0,
        "changes": 0,
        "control_snapshots": 0,
        "control_lines": 0,
    }
    for r in _rows(wb["PROJECTS"]):
        code = str(r.get("Project_ID") or "").strip()
        if not code:
            continue
        p = db.query(Project).filter(Project.code == code).first()
        if not p:
            p = Project(
                code=code,
                customer=str(r.get("Customer") or "Cliente importato"),
                sales_owner=str(r.get("Sales") or "N/D"),
                pm_owner=str(r.get("PM") or "N/D"),
                order_value=float(r.get("Order_Value") or 0),
                status=str(r.get("Status") or "EXECUTION"),
                technical_status=str(r.get("Technical_Status") or "IN_PROGRESS"),
                country=str(r.get("Country") or "IT"),
            )
            db.add(p)
            db.flush()
            stats["projects"] += 1
        else:
            p.customer = str(r.get("Customer") or p.customer)
            p.sales_owner = str(r.get("Sales") or p.sales_owner)
            p.pm_owner = str(r.get("PM") or p.pm_owner)
            p.order_value = float(r.get("Order_Value") or p.order_value)
            p.status = str(r.get("Status") or p.status)
    db.commit()
    idx = {p.code: p for p in db.query(Project).all()}

    control_idx = {}
    if "CONTROL_SNAPSHOT" in wb.sheetnames:
        for r in _rows(wb["CONTROL_SNAPSHOT"]):
            code = str(r.get("Project_ID") or "").strip()
            p = idx.get(code)
            if not p:
                continue
            snap_date = _d(r.get("Snapshot_Date")) or datetime.today().date()
            captured = datetime.combine(snap_date, datetime.min.time())
            source = str(r.get("Source_Label") or "GALILEO")
            snap = (
                db.query(ProjectControlSnapshot)
                .filter(
                    ProjectControlSnapshot.project_id == p.id,
                    ProjectControlSnapshot.captured_at == captured,
                    ProjectControlSnapshot.source_label == source,
                )
                .first()
            )
            if not snap:
                snap = ProjectControlSnapshot(project_id=p.id, captured_at=captured, source_label=source)
                db.add(snap)
                db.flush()
                stats["control_snapshots"] += 1
            snap.baseline_budget = float(r.get("Baseline_Budget") or 0)
            snap.operational_budget = float(r.get("Operational_Budget") or 0)
            snap.orders_jmrp = float(r.get("Orders_JMRP") or 0)
            snap.orders_extract = float(r.get("Orders_Extract") or 0)
            snap.received = float(r.get("Received") or 0)
            snap.open_order_commitment = float(r.get("Open_Order_Commitment") or 0)
            snap.labor = float(r.get("Labor") or 0)
            snap.overhead = float(r.get("Overhead") or 0)
            snap.note = str(r.get("Note") or "") or None
            control_idx[(code, snap_date.isoformat())] = snap
    if "CONTROL_LINES" in wb.sheetnames:
        # Replace line detail for snapshots present in this import; keeps snapshots append-only over time.
        cleared = set()
        for r in _rows(wb["CONTROL_LINES"]):
            code = str(r.get("Project_ID") or "").strip()
            p = idx.get(code)
            if not p:
                continue
            snap_date = _d(r.get("Snapshot_Date")) or datetime.today().date()
            snap = control_idx.get((code, snap_date.isoformat()))
            if not snap:
                continue
            if snap.id not in cleared:
                db.query(ProjectControlLine).filter(ProjectControlLine.snapshot_id == snap.id).delete(
                    synchronize_session=False
                )
                cleared.add(snap.id)
            raw_closed = r.get("Purchases_Closed")
            if raw_closed is None or str(raw_closed).strip() == "":
                closed = None
            else:
                closed = str(raw_closed).strip().lower() in ("1", "true", "yes", "si", "sì", "s")
            db.add(
                ProjectControlLine(
                    snapshot_id=snap.id,
                    vda_code=str(r.get("VDA_Code") or "N/D"),
                    vda_name=str(r.get("VDA_Name") or "Voce di analisi"),
                    analysis_center=str(r.get("Analysis_Center") or "COMM"),
                    baseline_budget=float(r.get("Baseline_Budget") or 0),
                    operational_budget=float(r.get("Operational_Budget") or 0),
                    ordered=float(r.get("Ordered") or 0),
                    received=float(r.get("Received") or 0),
                    open_order_commitment=float(r.get("Open_Order_Commitment") or 0),
                    labor=float(r.get("Labor") or 0),
                    overhead=float(r.get("Overhead") or 0),
                    purchases_closed=closed,
                    source_ref=str(r.get("Source_Ref") or "") or None,
                )
            )
            stats["control_lines"] += 1

    if "CUSTOMER_ORDERS" in wb.sheetnames:
        for r in _rows(wb["CUSTOMER_ORDERS"]):
            p = idx.get(str(r.get("Project_ID") or "").strip())
            if not p:
                continue
            number = str(r.get("Order_Number") or "").strip()
            if not number:
                continue
            existing = next((o for o in p.customer_orders if o.order_number == number), None)
            o = existing or CustomerOrder(
                order_number=number, order_date=_d(r.get("Order_Date")) or datetime.today().date()
            )
            o.order_date = _d(r.get("Order_Date")) or o.order_date
            o.order_kind = str(r.get("Order_Kind") or "BASE").upper()
            o.quote_ref = str(r.get("Quote_Ref") or "") or None
            o.net_value = float(r.get("Net_Value") or 0)
            o.tax_value = float(r.get("Tax_Value") or 0)
            o.gross_value = float(r.get("Gross_Value") or 0)
            o.currency = str(r.get("Currency") or "EUR")
            o.payment_terms = str(r.get("Payment_Terms") or "") or None
            o.incoterm = str(r.get("Incoterm") or "") or None
            o.scope_summary = str(r.get("Scope_Summary") or "") or None
            o.source_ref = str(r.get("Source_Ref") or "") or None
            o.note = str(r.get("Note") or "") or None
            if not existing:
                p.customer_orders.append(o)
                stats["customer_orders"] += 1
    if "CHANGES" in wb.sheetnames:
        for r in _rows(wb["CHANGES"]):
            p = idx.get(str(r.get("Project_ID") or "").strip())
            if not p:
                continue
            p.changes.append(
                ChangeRequest(
                    description=str(r.get("Description") or "Variante"),
                    requested_date=_d(r.get("Date")) or datetime.today().date(),
                    estimated_cost=float(r.get("Estimated_Cost") or 0),
                    quoted_value=float(r.get("Quoted_Value") or 0),
                    order_received=str(r.get("Order_Received") or "").strip().lower()
                    in ("1", "true", "yes", "si", "sì", "s"),
                    work_started=str(r.get("Work_Started") or "").strip().lower()
                    in ("1", "true", "yes", "si", "sì", "s"),
                    source_ref=str(r.get("Source_Ref") or "") or None,
                )
            )
            stats["changes"] += 1

    if "MILESTONES" in wb.sheetnames:
        for r in _rows(wb["MILESTONES"]):
            p = idx.get(str(r.get("Project_ID") or "").strip())
            if not p:
                continue
            p.milestones.append(
                Milestone(
                    name=str(r.get("Milestone") or "Milestone"),
                    contract_date=_d(r.get("Contract_Date")),
                    forecast_date=_d(r.get("Forecast_Date")),
                    actual_date=_d(r.get("Actual_Date")),
                    invoice_value=float(r.get("Invoice_Value") or 0),
                    invoice_required=bool(r.get("Invoice_Required") or False),
                    ld_applicable=bool(r.get("LD_Applicable") or False),
                    ld_rate_weekly_pct=float(r.get("LD_Rate_Weekly_Pct") or 0),
                )
            )
            stats["milestones"] += 1
    if "BUDGET" in wb.sheetnames:
        for r in _rows(wb["BUDGET"]):
            p = idx.get(str(r.get("Project_ID") or "").strip())
            if not p:
                continue
            p.budget_lines.append(
                BudgetLine(
                    category=str(r.get("Category") or "General"),
                    budget=float(r.get("Budget") or 0),
                    actual=float(r.get("Actual") or 0),
                    committed=float(r.get("Committed") or 0),
                    etc=float(r.get("ETC") or 0),
                    etc_source=str(r.get("ETC_Source") or "MANUAL"),
                )
            )
            stats["budget"] += 1
    if "INVOICES" in wb.sheetnames:
        for r in _rows(wb["INVOICES"]):
            p = idx.get(str(r.get("Project_ID") or "").strip())
            if not p:
                continue
            p.invoices.append(
                Invoice(
                    number=str(r.get("Invoice_Number") or "N/D"),
                    milestone_name=r.get("Milestone"),
                    issue_date=_d(r.get("Issue_Date")),
                    due_date=_d(r.get("Due_Date")),
                    amount=float(r.get("Amount") or 0),
                    paid_amount=float(r.get("Paid_Amount") or 0),
                )
            )
            stats["invoices"] += 1
    if "PAYMENTS" in wb.sheetnames:
        for r in _rows(wb["PAYMENTS"]):
            p = idx.get(str(r.get("Project_ID") or "").strip())
            if not p:
                continue
            p.payments.append(
                Payment(
                    invoice_number=str(r.get("Invoice_Number") or "N/D"),
                    payment_date=_d(r.get("Payment_Date")),
                    amount=float(r.get("Amount") or 0),
                )
            )
            stats["payments"] += 1
    if "ACTIONS" in wb.sheetnames:
        for r in _rows(wb["ACTIONS"]):
            p = idx.get(str(r.get("Project_ID") or "").strip())
            if not p:
                continue
            p.actions.append(
                Action(
                    title=str(r.get("Title") or "Action"),
                    owner=str(r.get("Owner") or "N/D"),
                    due_date=_d(r.get("Due_Date")),
                    status=str(r.get("Status") or "OPEN"),
                    source="EXCEL",
                )
            )
            stats["actions"] += 1
    if "COMMUNICATIONS" in wb.sheetnames:
        for r in _rows(wb["COMMUNICATIONS"]):
            p = idx.get(str(r.get("Project_ID") or "").strip())
            if not p:
                continue
            p.communications.append(
                Communication(
                    date=_d(r.get("Date")) or datetime.today().date(),
                    direction=str(r.get("Direction") or "INBOUND").upper(),
                    sender=str(r.get("Sender") or "external@example.com"),
                    recipient=str(r.get("Recipient") or "brofind@example.com"),
                    subject=str(r.get("Subject") or "Communication"),
                    body=str(r.get("Body") or ""),
                    topic=str(r.get("Topic") or "Imported"),
                    responded=str(r.get("Responded") or "").strip().lower() in ("1", "true", "yes", "si", "sì"),
                    reminders_count=int(r.get("Reminders_Count") or 0),
                    source_ref=str(r.get("Source_Ref") or "excel:communication"),
                )
            )
            stats["communications"] += 1
    db.commit()
    return stats
