from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

import app.b1_control_adapter as adapter
from app.b1_control_adapter import B1ImportError, import_b1_control_workbook, is_b1_workbook_candidate
from app.control import control_summary
from app.db import Base, SessionLocal, engine
from app.main import app
from app.models import Invoice, Project, ProjectControlLine, ProjectControlSnapshot, User
from app.rules import finance_summary
from app.seed import seed


FIXTURE = Path(__file__).resolve().parents[1] / "demo" / "b1" / "assets" / "DEMO-B1-001_control.xlsx"
EVALUATION_DATE = date(2026, 9, 28)


def workbook_bytes(mutator=None):
    workbook = openpyxl.load_workbook(FIXTURE)
    if mutator:
        mutator(workbook)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.fixture()
def db():
    Base.metadata.create_all(bind=engine)
    session = SessionLocal()
    seed(session)
    if not session.query(Project).filter_by(code="DEMO-B1-001").first():
        session.add(Project(code="DEMO-B1-001", customer="Test", sales_owner="S. Riva", pm_owner="A. Greco", order_value=500000))
        session.commit()
    project = session.query(Project).filter_by(code="DEMO-B1-001").one()
    snapshot_ids = [snapshot.id for snapshot in session.query(ProjectControlSnapshot).filter_by(project_id=project.id)]
    if snapshot_ids:
        session.query(ProjectControlLine).filter(ProjectControlLine.snapshot_id.in_(snapshot_ids)).delete(synchronize_session=False)
    session.query(ProjectControlSnapshot).filter_by(project_id=project.id).delete(synchronize_session=False)
    session.query(Invoice).filter_by(project_id=project.id).delete(synchronize_session=False)
    session.commit()
    try:
        yield session
    finally:
        session.close()


def counts(db):
    project = db.query(Project).filter_by(code="DEMO-B1-001").one()
    snapshots = db.query(ProjectControlSnapshot).filter_by(project_id=project.id).count()
    lines = db.query(ProjectControlLine).join(ProjectControlSnapshot).filter(ProjectControlSnapshot.project_id == project.id).count()
    invoices = db.query(Invoice).filter_by(project_id=project.id).count()
    return snapshots, lines, invoices


def test_b1_frozen_economic_contract_and_cda_aggregation(db):
    result = import_b1_control_workbook(FIXTURE.read_bytes(), db)
    assert result == {"control_snapshots": 1, "control_lines": 5, "invoices": 1}
    project = db.query(Project).filter_by(code="DEMO-B1-001").one()
    control = control_summary(project)
    assert (control["baseline_budget"], control["operational_budget"], control["orders_extract"]) == (340000, 360000, 320000)
    assert (control["received"], control["open_order_commitment"], control["labor"]) == (202000, 118000, 35000)
    assert (control["actual_total"], control["absorbed_committed"], control["uncommitted_budget"]) == (237000, 355000, 5000)
    carpenterie = next(row for row in control["vda_rows"] if row["vda_code"] == "C000013")
    elettrici = next(row for row in control["vda_rows"] if row["vda_code"] == "C000026")
    assert (carpenterie["operational_budget"], carpenterie["ordered"]) == (150000, 190000)
    assert control["budget_health"]["gross_overrun"] == 40000
    assert control["budget_health"]["uncovered_overrun"] == 40000
    assert (elettrici["operational_budget"], elettrici["ordered"], len(elettrici["centers"])) == (100000, 100000, 2)
    invoice = project.invoices[0]
    assert (invoice.number, invoice.amount, invoice.paid_amount, invoice.amount - invoice.paid_amount) == ("INV-DEMO-001", 250000, 50000, 200000)
    assert (EVALUATION_DATE - invoice.due_date).days == 28
    assert finance_summary(project, EVALUATION_DATE)["overdue"] == 200000


def test_b1_double_import_is_idempotent(db):
    content = FIXTURE.read_bytes()
    import_b1_control_workbook(content, db)
    before = control_summary(db.query(Project).filter_by(code="DEMO-B1-001").one())
    result = import_b1_control_workbook(content, db)
    after = control_summary(db.query(Project).filter_by(code="DEMO-B1-001").one())
    assert result == {"control_snapshots": 0, "control_lines": 5, "invoices": 0}
    assert counts(db) == (1, 5, 1)
    for key in ("baseline_budget", "operational_budget", "orders_extract", "actual_total", "absorbed_committed", "uncommitted_budget"):
        assert after[key] == before[key]


@pytest.mark.parametrize("mutator, message", [
    (lambda wb: wb.create_sheet("PAYMENTS"), "fogli non conformi"),
    (lambda wb: setattr(wb["CONTROL_LINES"][1][0], "value", "Wrong"), "header non conforme"),
    (lambda wb: setattr(wb["CONTROL_SNAPSHOT"]["D2"], "value", "=1+1"), "formula non ammessa"),
    (lambda wb: setattr(wb["CONTROL_LINES"]["M2"], "value", "maybe"), "TRUE, FALSE"),
    (lambda wb: setattr(wb["CONTROL_SNAPSHOT"]["H2"], "value", 1), "quadratura"),
])
def test_b1_invalid_temporary_workbooks_are_rejected_without_persistence(db, mutator, message):
    content = workbook_bytes(mutator)
    assert is_b1_workbook_candidate(content) is True
    with pytest.raises(B1ImportError, match=message):
        import_b1_control_workbook(content, db)
    assert counts(db) == (0, 0, 0)


def test_extra_sheet_is_b1_routed_rejected_and_never_falls_back_to_generic(monkeypatch):
    content = workbook_bytes(lambda wb: wb.create_sheet("EXTRA"))
    assert is_b1_workbook_candidate(content) is True
    Base.metadata.create_all(bind=engine)
    with SessionLocal() as db:
        seed(db)
        if not db.query(Project).filter_by(code="DEMO-B1-001").first():
            db.add(Project(code="DEMO-B1-001", customer="Test", sales_owner="S. Riva", pm_owner="A. Greco", order_value=500000))
            db.commit()
    with TestClient(app) as client:
        def generic_must_not_run(*args, **kwargs):
            raise AssertionError("generic importer invoked")
        monkeypatch.setattr("app.main.import_workbook", generic_must_not_run)
        response = client.post("/import", files={"file": ("bad-b1.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert response.status_code == 200
        assert "B1: fogli non conformi" in response.text
        with SessionLocal() as db:
            assert counts(db) == (0, 0, 0)


def test_duplicate_db_keys_fail_closed(db):
    project = db.query(Project).filter_by(code="DEMO-B1-001").one()
    db.add_all([
        Invoice(project_id=project.id, number="INV-DEMO-001", milestone_name="x", issue_date=date(2026, 8, 1), due_date=date(2026, 8, 31), amount=1, paid_amount=0),
        Invoice(project_id=project.id, number="INV-DEMO-001", milestone_name="x", issue_date=date(2026, 8, 1), due_date=date(2026, 8, 31), amount=1, paid_amount=0),
    ])
    db.commit()
    with pytest.raises(B1ImportError, match="duplicati gia presenti nel DB"):
        import_b1_control_workbook(FIXTURE.read_bytes(), db)
    assert counts(db) == (0, 0, 2)


def test_forced_failure_after_control_write_rolls_back(db, monkeypatch):
    def fail_after_control_write(*args, **kwargs):
        raise RuntimeError("forced invoice persistence failure")
    monkeypatch.setattr(adapter, "_upsert_invoice", fail_after_control_write)
    with pytest.raises(RuntimeError, match="forced invoice"):
        import_b1_control_workbook(FIXTURE.read_bytes(), db)
    assert counts(db) == (0, 0, 0)

def test_duplicate_snapshot_db_key_fails_closed(db):
    project = db.query(Project).filter_by(code="DEMO-B1-001").one()
    source = adapter._parse(FIXTURE.read_bytes()).snapshot
    captured = datetime.combine(source["snapshot_date"], datetime.min.time())
    db.add_all([
        ProjectControlSnapshot(project_id=project.id, captured_at=captured, source_label=source["source_label"]),
        ProjectControlSnapshot(project_id=project.id, captured_at=captured, source_label=source["source_label"]),
    ])
    db.commit()
    with pytest.raises(B1ImportError, match="duplicati gia presenti nel DB"):
        import_b1_control_workbook(FIXTURE.read_bytes(), db)
    assert counts(db) == (2, 0, 0)